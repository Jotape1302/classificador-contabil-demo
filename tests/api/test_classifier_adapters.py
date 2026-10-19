"""Paridade dos adaptadores com DataFrames e regras legadas."""

import datetime
import io
from types import SimpleNamespace

from openpyxl import load_workbook
import pandas as pd

from api.zaeron_api.adapters.parsers import (
    ler_extrato_csv,
    ler_extrato_ofx,
)
from api.zaeron_api.services.classifier_engine import (
    classificar_transacoes_hibrido,
)
from api.zaeron_api.adapters import parser_adapter
from api.zaeron_api.adapters.classifier_adapter import classify_transactions
from api.zaeron_api.adapters.export_adapter import build_export


OFX_BYTES = b"""OFXHEADER:100
DATA:OFXSGML
VERSION:102
<OFX><BANKMSGSRSV1><STMTTRNRS><STMTRS><BANKTRANLIST>
<STMTTRN><DTPOSTED>20260810120000<TRNAMT>-450.00<FITID>1<MEMO>COPEL ENERGIA ELETRICA</STMTTRN>
<STMTTRN><DTPOSTED>20260811120000<TRNAMT>-68.90<FITID>2<MEMO>TARIFA BANCARIA PACOTE</STMTTRN>
</BANKTRANLIST></STMTRS></STMTTRNRS></BANKMSGSRSV1></OFX>"""


def test_ofx_adapter_has_parity_with_legacy_dataframe():
    legacy_df, error = ler_extrato_ofx(OFX_BYTES, user_id=None)
    parsed = parser_adapter.parse_statement("extrato.ofx", OFX_BYTES, user_id=7)

    assert error is None
    assert len(parsed.transactions) == len(legacy_df) == 2
    assert parsed.transactions[0]["date"].isoformat() == legacy_df.iloc[0]["Data"]
    assert parsed.transactions[0]["clean_description"] == legacy_df.iloc[0]["Descricao_Limpa"]
    assert parsed.transactions[0]["amount"] == legacy_df.iloc[0]["Valor"]


def test_csv_adapter_has_parity_with_legacy_dataframe():
    csv_bytes = "Data;Histórico;Valor\n10/08/2026;COPEL ENERGIA;-450,00\n".encode(
        "utf-8"
    )
    legacy_df, error = ler_extrato_csv(csv_bytes, user_id=None)
    parsed = parser_adapter.parse_statement("extrato.csv", csv_bytes, user_id=9)

    assert error is None
    assert len(parsed.transactions) == len(legacy_df) == 1
    assert parsed.transactions[0]["original_description"] == legacy_df.iloc[0]["Descricao_Original"]
    assert parsed.transactions[0]["amount"] == -450.0


def test_pdf_adapter_preserves_the_same_legacy_dataframe(monkeypatch):
    legacy_df = pd.DataFrame(
        [
            {
                "Data": "2026-08-12",
                "Descricao_Original": "SANEPAR AGUA",
                "Descricao_Limpa": "SANEPAR AGUA",
                "Valor": -120.0,
                "Tipo": "Débito",
                "Categoria": "Pendente",
                "Conta_Debito": "",
                "Conta_Credito": "",
                "Historico": "Vlr. ref. pagamento SANEPAR AGUA",
                "Origem": "Não classificado",
                "Confianca": 0.0,
                "Is_Duplicado": False,
            }
        ]
    )
    monkeypatch.setitem(
        parser_adapter._PARSERS,
        "pdf",
        lambda _content, user_id=None: (legacy_df.copy(), None),
    )

    parsed = parser_adapter.parse_statement("extrato.pdf", b"%PDF", user_id=11)

    assert parsed.transactions[0]["accounting_history"] == legacy_df.iloc[0]["Historico"]
    assert parsed.transactions[0]["history_code"] is None


def test_pure_classifier_keeps_legacy_results_without_database_effects():
    transactions = [
        SimpleNamespace(
            id=1,
            original_description="COPEL ENERGIA ELETRICA",
            clean_description="COPEL ENERGIA ELETRICA",
            amount=-450.0,
            is_duplicate=False,
        )
    ]
    adapted, error, stats, rule_ids = classify_transactions(transactions, rules=[])
    legacy, legacy_error, legacy_stats = classificar_transacoes_hibrido(
        [
            {
                "id": 1,
                "Descricao_Original": "COPEL ENERGIA ELETRICA",
                "Descricao_Limpa": "COPEL ENERGIA ELETRICA",
                "Valor": -450.0,
            }
        ]
    )

    assert error == legacy_error == ""
    assert adapted[0]["category"] == legacy[0]["categoria"] == "Energia Elétrica"
    assert adapted[0]["confidence"] == legacy[0]["confianca"]
    assert stats == legacy_stats
    assert rule_ids == []


def test_export_keeps_narrative_history_separate_and_sanitizes_html():
    transaction = SimpleNamespace(
        date=datetime.date(2026, 8, 10),
        original_description="<script>alert(1)</script>",
        clean_description="<script>alert(1)</script>",
        amount=-100.0,
        transaction_type="Débito",
        category="Serviços de Terceiros",
        debit_account="3.1.1",
        credit_account="1.1.1",
        accounting_history="Vlr. ref. pagamento de serviço",
        history_code="77",
        classification_source="Edição manual",
        confidence=1,
        is_duplicate=False,
    )

    dominio = build_export(
        "dominio", [transaction], office_name="Zaeron", accountant_name="Ana"
    ).content.decode("utf-8-sig")
    report = build_export(
        "html", [transaction], office_name="<b>Office</b>", accountant_name="Ana"
    ).content.decode("utf-8")

    assert ";77;Vlr. ref. pagamento de serviÃ§o" not in dominio
    assert ";77;Vlr. ref. pagamento de serviço" in dominio
    assert "<script>alert(1)</script>" not in report
    assert "&lt;b&gt;Office&lt;/b&gt;" in report


def test_excel_export_has_five_columns_without_debit_credit():
    transaction = SimpleNamespace(
        date=datetime.date(2026, 8, 10),
        original_description="COPEL ENERGIA ELETRICA",
        clean_description="COPEL ENERGIA ELETRICA",
        amount=-450.0,
        transaction_type="Débito",
        category="Energia Elétrica",
        debit_account="3.1.1.01",
        credit_account="1.1.1.01",
        accounting_history="Vlr. ref. pagamento COPEL ENERGIA",
        history_code="10",
        classification_source="Regra do Escritório",
        confidence=1.0,
        is_duplicate=False,
    )

    exported = build_export(
        "xlsx", [transaction], office_name="Zaeron Tech", accountant_name="Carlos"
    )
    wb = load_workbook(io.BytesIO(exported.content))
    ws = wb["Conciliação Zaeron"]

    headers = [cell.value for cell in ws[4] if cell.value is not None]
    assert headers == ["Data", "Histórico Contábil", "Categoria", "Valor (R$)", "Status"]
    assert "Conta Débito" not in headers
    assert "Conta Crédito" not in headers
    assert "Débito" not in headers
    assert "Crédito" not in headers

    row_vals = [cell.value for cell in ws[5] if cell.value is not None]
    assert row_vals == ["10/08/2026", "Vlr. ref. pagamento COPEL ENERGIA", "Energia Elétrica", -450.0, "Conciliado"]


def test_pdf_export_generates_five_column_report_without_debit_credit():
    transaction = SimpleNamespace(
        date=datetime.date(2026, 8, 10),
        original_description="SANEPAR AGUA",
        clean_description="SANEPAR AGUA",
        amount=-120.0,
        transaction_type="Débito",
        category="Água e Esgoto",
        debit_account="3.1.1.02",
        credit_account="1.1.1.01",
        accounting_history="Vlr. ref. pagamento SANEPAR AGUA",
        history_code="10",
        classification_source="IA Gemini",
        confidence=0.95,
        is_duplicate=False,
    )

    exported = build_export(
        "pdf", [transaction], office_name="Zaeron Tech", accountant_name="Carlos"
    )
    assert exported.media_type == "application/pdf"
    assert exported.extension == "pdf"
    assert exported.content.startswith(b"%PDF")
    assert len(exported.content) > 1000


def test_csv_export_has_five_columns_without_debit_credit():
    transaction = SimpleNamespace(
        date=datetime.date(2026, 8, 10),
        original_description="COPEL ENERGIA",
        clean_description="COPEL ENERGIA",
        amount=-450.0,
        transaction_type="Débito",
        category="Energia Elétrica",
        debit_account="3.1.1.01",
        credit_account="1.1.1.01",
        accounting_history="Vlr. ref. pagamento COPEL ENERGIA",
        history_code="10",
        classification_source="Regra do Escritório",
        confidence=1.0,
        is_duplicate=False,
    )

    exported = build_export(
        "csv", [transaction], office_name="Zaeron Tech", accountant_name="Carlos"
    )
    lines = exported.content.decode("utf-8-sig").splitlines()
    assert lines[0] == "Data;Histórico Contábil;Categoria;Valor;Status"
    assert "Conta_Debito" not in lines[0]
    assert "Conta_Credito" not in lines[0]
    assert lines[1] == "10/08/2026;Vlr. ref. pagamento COPEL ENERGIA;Energia Elétrica;-450.0;Conciliado"


def test_all_eight_export_formats_and_sanitization():
    from api.zaeron_api.routers.classifier import _build_safe_export_filename

    transaction = SimpleNamespace(
        date=datetime.date(2026, 8, 10),
        original_description="PAGAMENTO COPEL ENERGIA",
        clean_description="COPEL ENERGIA",
        amount=-350.50,
        transaction_type="Débito",
        category="Energia Elétrica",
        debit_account="3.1.1.01",
        credit_account="1.1.1.01",
        accounting_history="Vlr. ref. pagamento COPEL ENERGIA",
        history_code="10",
        classification_source="IA Gemini",
        confidence=0.92,
        is_duplicate=False,
    )
    transactions = [transaction]

    # 1. Domínio (TXT)
    exp_dom = build_export("dominio", transactions, office_name="Zaeron Tech", accountant_name="Carlos")
    assert exp_dom.extension == "txt"
    assert exp_dom.media_type.startswith("text/plain")
    assert "10/08/2026;3.1.1.01;1.1.1.01;350,50;10;Vlr. ref. pagamento COPEL ENERGIA" in exp_dom.content.decode("utf-8-sig")

    # 2. Questor (TXT)
    exp_que = build_export("questor", transactions, office_name="Zaeron Tech", accountant_name="Carlos")
    assert exp_que.extension == "txt"
    assert exp_que.media_type.startswith("text/plain")
    assert "10/08/2026;3.1.1.01;1.1.1.01;350,50;Vlr. ref. pagamento COPEL ENERGIA;1;10" in exp_que.content.decode("utf-8-sig")

    # 3. Omie (CSV)
    exp_omi = build_export("omie", transactions, office_name="Zaeron Tech", accountant_name="Carlos")
    assert exp_omi.extension == "csv"
    assert exp_omi.media_type.startswith("text/csv")
    assert "10/08/2026;350,50;Energia Elétrica;Vlr. ref. pagamento COPEL ENERGIA;CONTA CORRENTE;Débito" in exp_omi.content.decode("utf-8-sig")

    # 4. Fortes / Contmatic (CSV)
    exp_for = build_export("fortes", transactions, office_name="Zaeron Tech", accountant_name="Carlos")
    assert exp_for.extension == "csv"
    assert exp_for.media_type.startswith("text/csv")
    fortes_text = exp_for.content.decode("utf-8-sig")
    assert "Origem_Classificacao" not in fortes_text
    assert "IA Gemini" not in fortes_text
    assert "10/08/2026;3.1.1.01;1.1.1.01;350,50;Vlr. ref. pagamento COPEL ENERGIA;Energia Elétrica;Conciliado" in fortes_text

    # 5. XLSX (Excel)
    exp_xlsx = build_export("xlsx", transactions, office_name="Zaeron Tech", accountant_name="Carlos")
    assert exp_xlsx.extension == "xlsx"
    assert "openxmlformats" in exp_xlsx.media_type
    assert len(exp_xlsx.content) > 500

    # 6. CSV Padrão
    exp_csv = build_export("csv", transactions, office_name="Zaeron Tech", accountant_name="Carlos")
    assert exp_csv.extension == "csv"
    assert "10/08/2026;Vlr. ref. pagamento COPEL ENERGIA;Energia Elétrica;-350.5;Conciliado" in exp_csv.content.decode("utf-8-sig")

    # 7. PDF Executivo
    exp_pdf = build_export("pdf", transactions, office_name="Zaeron Tech", accountant_name="Carlos")
    assert exp_pdf.extension == "pdf"
    assert exp_pdf.content.startswith(b"%PDF")

    # 8. HTML Executivo
    exp_html = build_export("html", transactions, office_name="Zaeron Tech", accountant_name="Carlos", source_filename="extrato_agosto_2026.ofx", batch_id=42)
    assert exp_html.extension == "html"
    html_text = exp_html.content.decode("utf-8")
    assert "ZAERON CONTÁBIL" in html_text
    assert "Relatório Executivo de Conciliação" in html_text
    assert "extrato_agosto_2026.ofx" in html_text
    assert "Lote #42" in html_text
    assert "R$ 11,00/h" in html_text
    assert "IA Gemini" not in html_text
    assert "Regra Nativa" not in html_text

    # Filename generation tests
    fn1 = _build_safe_export_filename("Extrato Itaú 08-2026.OFX", 42, "xlsx", "xlsx")
    assert fn1 == "classificacao_extrato_itau_08-2026.xlsx"
    assert "anonymous" not in fn1

    fn2 = _build_safe_export_filename(None, 42, "pdf", "pdf")
    assert fn2 == "zaeron_lote_42_pdf.pdf"
    assert "anonymous" not in fn2


