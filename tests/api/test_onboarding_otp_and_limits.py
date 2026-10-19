"""Testes para o fluxo de cadastro direto, validações estritas de CPF/CNPJ, prevenção de duplicidade, limites e exportação."""

import datetime
from io import BytesIO
import openpyxl
from pypdf import PdfReader
import pytest

from api.zaeron_api.models.users import User
from api.zaeron_api.core.validators import (
    clean_digits,
    format_cep,
    validate_cep,
    validate_cnpj,
    validate_cpf,
    validate_document,
)


# ==============================================================================
# 1. TESTES DE VALIDAÇÃO DE DOCUMENTOS (CPF / CNPJ)
# ==============================================================================

def test_cpf_validation():
    # CPFs válidos
    assert validate_cpf("52998224725") is True
    assert validate_cpf("529.982.247-25") is True
    assert validate_cpf("12345678909") is True

    # CPFs inválidos: dígitos repetidos
    assert validate_cpf("11111111111") is False
    assert validate_cpf("00000000000") is False
    assert validate_cpf("99999999999") is False

    # CPFs inválidos: dígitos verificadores incorretos
    assert validate_cpf("52998224726") is False
    assert validate_cpf("12345678900") is False

    # CPFs com tamanho incorreto
    assert validate_cpf("12345678") is False
    assert validate_cpf("123456789012") is False


def test_cnpj_validation():
    # CNPJs válidos
    assert validate_cnpj("11222333000181") is True
    assert validate_cnpj("11.222.333/0001-81") is True
    assert validate_cnpj("00000000000191") is True

    # CNPJs inválidos: dígitos repetidos
    assert validate_cnpj("11111111111111") is False
    assert validate_cnpj("00000000000000") is False

    # CNPJs inválidos: dígitos verificadores errados
    assert validate_cnpj("11222333000182") is False

    # CNPJs com tamanho incorreto
    assert validate_cnpj("11222333") is False


def test_validate_document_helper():
    valid_cpf, digits, doc_type = validate_document("529.982.247-25")
    assert valid_cpf is True
    assert digits == "52998224725"
    assert doc_type == "cpf"

    valid_cnpj, digits, doc_type = validate_document("11.222.333/0001-81")
    assert valid_cnpj is True
    assert digits == "11222333000181"
    assert doc_type == "cnpj"

    invalid, _, _ = validate_document("123")
    assert invalid is False

    invalid_cpf_repeat, _, _ = validate_document("111.111.111-11")
    assert invalid_cpf_repeat is False


# ==============================================================================
# 2. TESTES DE CADASTRO DIRETO & SESSÃO IMEDIATA (SEM OTP, SEM ENDEREÇO)
# ==============================================================================

def test_direct_registration_success_and_immediate_login(client, session_factory):
    # 1. Cadastro com payload direto (nome, e-mail, senha, CPF)
    resp = client.post(
        "/api/auth/register",
        json={
            "name": "Carlos Eduardo Contador",
            "email": "carlos.eduardo@contabilidade.test",
            "password": "minha-senha-forte-123",
            "document_number": "529.982.247-25",
        },
    )
    assert resp.status_code == 201
    data = resp.json()
    assert data["user"]["email"] == "carlos.eduardo@contabilidade.test"
    assert data["user"]["name"] == "Carlos Eduardo Contador"
    assert data["user"]["document_number"] == "52998224725"
    assert data["user"]["is_verified"] is True
    assert "password_hash" not in resp.text
    assert "zaeron_session=" in resp.headers.get("set-cookie", "")

    # 2. Verifica no banco que o usuário já nasce verificado
    session = session_factory()
    try:
        user = session.query(User).filter_by(email="carlos.eduardo@contabilidade.test").one()
        assert user.is_verified is True
        assert user.document_number == "52998224725"
        assert user.plan_tier == "free"
        assert user.monthly_usage_count == 0
    finally:
        session.close()

    # 3. Usuário já tem sessão ativa e pode acessar rotas autenticadas imediatamente
    me_resp = client.get("/api/auth/me")
    assert me_resp.status_code == 200
    assert me_resp.json()["user"]["email"] == "carlos.eduardo@contabilidade.test"


def test_direct_registration_with_cnpj(client):
    resp = client.post(
        "/api/auth/register",
        json={
            "name": "Silva & Associados Contabilidade",
            "email": "contato@silvacontabil.test",
            "password": "senha-segura-123",
            "document_number": "11.222.333/0001-81",
            "office_name": "Silva Contabilidade",
        },
    )
    assert resp.status_code == 201
    assert resp.json()["user"]["document_number"] == "11222333000181"
    assert resp.json()["user"]["is_verified"] is True


def test_direct_registration_rejects_invalid_documents(client):
    # CPF falso / repetido
    bad_cpf = client.post(
        "/api/auth/register",
        json={
            "name": "Contador Falso",
            "email": "falso@zaeron.test",
            "password": "senha-segura-123",
            "document_number": "111.111.111-11",
        },
    )
    assert bad_cpf.status_code == 422
    assert "CPF ou CNPJ inválido." in bad_cpf.text

    # CNPJ falso / repetido
    bad_cnpj = client.post(
        "/api/auth/register",
        json={
            "name": "Empresa Falsa",
            "email": "falsa@zaeron.test",
            "password": "senha-segura-123",
            "document_number": "00.000.000/0000-00",
        },
    )
    assert bad_cnpj.status_code == 422
    assert "CPF ou CNPJ inválido." in bad_cnpj.text

    # Dígitos verificadores incorretos
    bad_digits = client.post(
        "/api/auth/register",
        json={
            "name": "Contador Invalido",
            "email": "digitos@zaeron.test",
            "password": "senha-segura-123",
            "document_number": "529.982.247-26",
        },
    )
    assert bad_digits.status_code == 422
    assert "CPF ou CNPJ inválido." in bad_digits.text


def test_document_and_email_conflict_rejection(client):
    # 1. Cadastra primeiro usuário com CPF
    first = client.post(
        "/api/auth/register",
        json={
            "name": "Primeiro Usuário",
            "email": "primeiro@zaeron.test",
            "password": "senha-segura-123",
            "document_number": "123.456.789-09",
        },
    )
    assert first.status_code == 201

    # 2. Tentativa com mesmo e-mail (duplicidade de e-mail -> 409)
    email_conflict = client.post(
        "/api/auth/register",
        json={
            "name": "Segundo Usuário",
            "email": "primeiro@zaeron.test",
            "password": "senha-segura-123",
            "document_number": "529.982.247-25",
        },
    )
    assert email_conflict.status_code == 409
    assert email_conflict.json()["code"] == "email_already_registered"
    assert "Este e-mail já está cadastrado." in email_conflict.json()["detail"]

    # 3. Tentativa com mesmo documento (duplicidade de documento -> 409)
    doc_conflict = client.post(
        "/api/auth/register",
        json={
            "name": "Terceiro Usuário",
            "email": "terceiro@zaeron.test",
            "password": "senha-segura-123",
            "document_number": "123.456.789-09",
        },
    )
    assert doc_conflict.status_code == 409
    assert doc_conflict.json()["code"] == "document_already_registered"
    assert "Este CPF/CNPJ já está cadastrado." in doc_conflict.json()["detail"]


# ==============================================================================
# 3. TESTES DE USO ILIMITADO (SEM TRAVA DE 5 USOS)
# ==============================================================================

def test_unlimited_usage_processing(client, legacy_user, session_factory):
    # Login
    client.post(
        "/api/auth/login",
        json={"email": legacy_user.email, "password": "senha-legada"},
    )

    # Cria arquivo OFX de teste
    ofx_content = b"""OFXHEADER:100
DATA:OFXSGML
<OFX>
<BANKMSGSRSV1><STMTTRNRS><STMTRS>
<BANKTRANLIST>
<STMTTRN>
<TRNTYPE>DEBIT
<DTPOSTED>20260810
<TRNAMT>-150.00
<FITID>TX001
<MEMO>PAGAMENTO ENERGIA COPEL
</STMTTRN>
</BANKTRANLIST>
</STMTRS></STMTTRNRS></BANKMSGSRSV1>
</OFX>"""

    # Realiza upload de 6 lotes diferentes
    batch_ids = []
    for i in range(6):
        content = ofx_content.replace(b"TX001", f"TX00{i}".encode())
        upload_resp = client.post(
            "/api/classifier/upload",
            files={"file": (f"extrato_{i}.ofx", content, "application/octet-stream")},
        )
        assert upload_resp.status_code == 201
        batch_ids.append(upload_resp.json()["id"])

    # Processa os 6 lotes continuamente sem bloqueio de 403 Forbidden
    for i in range(6):
        proc_resp = client.post(
            "/api/classifier/process",
            json={"batch_id": batch_ids[i]},
        )
        assert proc_resp.status_code == 200

    # Verifica dados no /api/auth/me (tier ilimitado)
    me_resp = client.get("/api/auth/me")
    assert me_resp.status_code == 200
    quotas = me_resp.json()["quotas"]
    assert quotas["monthly_usage_count"] == 6
    assert quotas["plan_tier"] == "unlimited"
    assert quotas["max_monthly_usages"] is None


# ==============================================================================
# 4. TESTES DE EXPORTAÇÃO PDF & EXCEL (XLSX)
# ==============================================================================

def test_export_xlsx_and_pdf_formats(client, legacy_user):
    client.post(
        "/api/auth/login",
        json={"email": legacy_user.email, "password": "senha-legada"},
    )

    ofx_content = b"""OFXHEADER:100
DATA:OFXSGML
<OFX>
<BANKMSGSRSV1><STMTTRNRS><STMTRS>
<BANKTRANLIST>
<STMTTRN>
<TRNTYPE>DEBIT
<DTPOSTED>20260815
<TRNAMT>-320.50
<FITID>TX_EXP_01
<MEMO>INTERNET FIBRA EMPRESA
</STMTTRN>
<STMTTRN>
<TRNTYPE>CREDIT
<DTPOSTED>20260816
<TRNAMT>1500.00
<FITID>TX_EXP_02
<MEMO>RECEBIMENTO HONORARIOS
</STMTTRN>
</BANKTRANLIST>
</STMTRS></STMTTRNRS></BANKMSGSRSV1>
</OFX>"""

    upload = client.post(
        "/api/classifier/upload",
        files={"file": ("extrato_export.ofx", ofx_content, "application/octet-stream")},
    )
    batch_id = upload.json()["id"]

    client.post("/api/classifier/process", json={"batch_id": batch_id})

    # 1. Exportação Excel XLSX via /api/classifier/export
    resp_xlsx = client.post(
        "/api/classifier/export",
        json={"batch_id": batch_id, "format": "xlsx"},
    )
    assert resp_xlsx.status_code == 200
    assert resp_xlsx.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert 'filename="classificacao_extrato_export.xlsx"' in resp_xlsx.headers["content-disposition"]

    # Valida estrutura do Excel gerado
    wb = openpyxl.load_workbook(BytesIO(resp_xlsx.content))
    assert "Conciliação Zaeron" in wb.sheetnames
    assert "Resumo por Categoria" in wb.sheetnames
    ws1 = wb["Conciliação Zaeron"]
    assert "ZAERON CONTÁBIL" in str(ws1["A1"].value)

    # 2. Exportação PDF via /api/classifier/export
    resp_pdf = client.post(
        "/api/classifier/export",
        json={"batch_id": batch_id, "format": "pdf"},
    )
    assert resp_pdf.status_code == 200
    assert resp_pdf.headers["content-type"] == "application/pdf"
    assert 'filename="classificacao_extrato_export.pdf"' in resp_pdf.headers["content-disposition"]

    # Valida bytes do PDF gerado
    pdf_reader = PdfReader(BytesIO(resp_pdf.content))
    assert len(pdf_reader.pages) >= 1
    page_text = pdf_reader.pages[0].extract_text()
    assert "ZAERON CONTÁBIL" in page_text or "Conciliação" in page_text

    # 3. Exportação direta via GET /api/classifier/export/excel
    resp_direct_excel = client.get(f"/api/classifier/export/excel?batch_id={batch_id}")
    assert resp_direct_excel.status_code == 200
    assert resp_direct_excel.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    # 4. Exportação direta via GET /api/classifier/export/pdf
    resp_direct_pdf = client.get(f"/api/classifier/export/pdf?batch_id={batch_id}")
    assert resp_direct_pdf.status_code == 200
    assert resp_direct_pdf.headers["content-type"] == "application/pdf"
