"""Converte o schema canônico para os layouts de exportação (Domínio, Questor, Omie, Fortes, XLSX, CSV, HTML, PDF)."""

import datetime
import html
import io
from dataclasses import dataclass
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import pandas as pd

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm, mm
from reportlab.pdfgen import canvas
from reportlab.platypus import HRFlowable, KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from api.zaeron_api.services.export_service import (
    gerar_layout_dominio,
    gerar_layout_fortes_contmatic,
    gerar_layout_omie,
    gerar_layout_questor,
    gerar_relatorio_gerencial_html,
)


@dataclass(frozen=True)
class ExportedFile:
    content: bytes
    media_type: str
    extension: str


def transactions_to_legacy_dataframe(transactions: list[Any]) -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "Data": transaction.date.isoformat() if hasattr(transaction.date, "isoformat") else str(transaction.date),
                "Descricao_Original": transaction.original_description,
                "Descricao_Limpa": transaction.clean_description,
                "Valor": float(transaction.amount),
                "Tipo": transaction.transaction_type,
                "Categoria": transaction.category,
                "Conta_Debito": transaction.debit_account or "",
                "Conta_Credito": transaction.credit_account or "",
                "Historico": transaction.accounting_history,
                "history_code": transaction.history_code or "",
                "Origem": transaction.classification_source,
                "Confianca": float(transaction.confidence),
                "Is_Duplicado": bool(transaction.is_duplicate),
            }
            for transaction in transactions
        ]
    )


def _safe_spreadsheet_text(value: Any) -> Any:
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _sanitized_html_dataframe(dataframe: pd.DataFrame) -> pd.DataFrame:
    safe = dataframe.copy()
    for column in safe.select_dtypes(include=["object", "string"]).columns:
        safe[column] = safe[column].map(lambda value: html.escape(str(value), quote=True))
    return safe


# ==============================================================================
# GERADOR EXCEL FORMATADO (.XLSX)
# ==============================================================================

def gerar_excel_zaeron(transactions: list[Any], office_name: str, accountant_name: str) -> bytes:
    wb = Workbook()
    
    # Aba 1: Lançamentos
    ws1 = wb.active
    ws1.title = "Conciliação Zaeron"
    ws1.views.sheetView[0].showGridLines = True

    # Estilos Zaeron
    header_fill = PatternFill(start_color="171B20", end_color="171B20", fill_type="solid")
    header_font = Font(name="Arial", size=10, bold=True, color="F4EFEA")
    title_font = Font(name="Arial", size=13, bold=True, color="C8753D")
    meta_font = Font(name="Arial", size=9, italic=True, color="64748B")
    data_font = Font(name="Arial", size=9, color="1E293B")
    bold_data_font = Font(name="Arial", size=9, bold=True, color="1E293B")
    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )

    # Cabeçalho institucional
    ws1["A1"] = "ZAERON CONTÁBIL — RELATÓRIO DE CONCILIAÇÃO"
    ws1["A1"].font = title_font
    ws1["A2"] = f"Escritório: {office_name} | Responsável: {accountant_name} | Gerado em: {datetime.datetime.now().strftime('%d/%m/%Y %H:%M')}"
    ws1["A2"].font = meta_font

    headers = [
        "Data",
        "Histórico Contábil",
        "Categoria",
        "Valor (R$)",
        "Status",
    ]

    header_row = 4
    for col_num, header in enumerate(headers, 1):
        cell = ws1.cell(row=header_row, column=col_num, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center" if col_num == 1 else "right" if col_num == 4 else "left", vertical="center")
        cell.border = thin_border

    current_row = 5
    for tx in transactions:
        # Formata data
        dt_val = tx.date.strftime("%d/%m/%Y") if hasattr(tx.date, "strftime") else str(tx.date)
        if "-" in dt_val and len(dt_val.split("-")) == 3:
            p = dt_val.split("-")
            dt_val = f"{p[2]}/{p[1]}/{p[0]}"

        amount_val = float(tx.amount)
        conf_val = float(getattr(tx, "confidence", 0) or 0)
        status_origem = "Conciliado" if conf_val >= 0.8 else "Pendente"

        row_values = [
            dt_val,
            tx.accounting_history or tx.clean_description or tx.original_description,
            tx.category or "Outras Despesas",
            amount_val,
            status_origem,
        ]

        for col_num, val in enumerate(row_values, 1):
            cell = ws1.cell(row=current_row, column=col_num, value=val)
            cell.font = data_font
            cell.border = thin_border
            if col_num == 1:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_num == 4:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = '"R$ "#,##0.00;[Red]"R$ "-#,##0.00;"R$ "0.00'
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
        current_row += 1

    # Aba 2: Resumo por Categoria
    ws2 = wb.create_sheet(title="Resumo por Categoria")
    ws2.views.sheetView[0].showGridLines = True
    ws2["A1"] = "RESUMO POR CATEGORIA CONTÁBIL"
    ws2["A1"].font = title_font

    cat_headers = ["Categoria Contábil", "Qtd Lançamentos", "Total Movimentado (R$)", "Participação (%)"]
    for col_num, header in enumerate(cat_headers, 1):
        cell = ws2.cell(row=3, column=col_num, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="left" if col_num == 1 else "right", vertical="center")
        cell.border = thin_border

    df = transactions_to_legacy_dataframe(transactions)
    if not df.empty:
        df["Valor_Abs"] = df["Valor"].abs()
        total_geral = df["Valor_Abs"].sum()
        resumo = df.groupby("Categoria").agg(
            Qtd=("Descricao_Limpa", "count"),
            Total=("Valor_Abs", "sum"),
        ).reset_index().sort_values(by="Total", ascending=False)

        cat_row = 4
        for _, r in resumo.iterrows():
            pct = (r["Total"] / total_geral * 100) if total_geral > 0 else 0.0
            row_vals = [r["Categoria"], int(r["Qtd"]), float(r["Total"]), float(pct)]
            for col_num, val in enumerate(row_vals, 1):
                cell = ws2.cell(row=cat_row, column=col_num, value=val)
                cell.font = data_font
                cell.border = thin_border
                if col_num == 1:
                    cell.alignment = Alignment(horizontal="left", vertical="center")
                elif col_num == 2:
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                    cell.number_format = "#,##0"
                elif col_num == 3:
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                    cell.number_format = '"R$ "#,##0.00'
                elif col_num == 4:
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                    cell.number_format = "0.00"
            cat_row += 1

    # Auto-ajuste de largura de colunas para ambas as abas
    for sheet in [ws1, ws2]:
        for col in sheet.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                if len(val_str) > max_len and cell.row > 2:
                    max_len = len(val_str)
            sheet.column_dimensions[col_letter].width = max(max_len + 4, 12)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ==============================================================================
# GERADOR PDF INSTITUCIONAL COM PAGINAÇÃO
# ==============================================================================

class NumberedCanvas(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_footer(num_pages)
            super().showPage()
        super().save()

    def draw_page_footer(self, page_count: int):
        self.saveState()
        self.setFont("Helvetica", 8)
        self.setFillColor(colors.HexColor("#64748B"))
        # Linha divisória de rodapé
        self.setStrokeColor(colors.HexColor("#CBD5E1"))
        self.setLineWidth(0.5)
        self.line(30, 30, self._pagesize[0] - 30, 30)

        # Texto do rodapé
        self.drawString(30, 20, "Zaeron Contábil • Plataforma de Inteligência e Automação Contábil")
        page_text = f"Página {self._pageNumber} de {page_count}"
        self.drawRightString(self._pagesize[0] - 30, 20, page_text)
        self.restoreState()


def gerar_pdf_zaeron(transactions: list[Any], office_name: str, accountant_name: str) -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=20 * mm,
        rightMargin=20 * mm,
        topMargin=15 * mm,
        bottomMargin=20 * mm,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "ReportTitle",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=16,
        leading=20,
        textColor=colors.HexColor("#0F172A"),
    )
    subtitle_style = ParagraphStyle(
        "ReportSubtitle",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9,
        leading=12,
        textColor=colors.HexColor("#64748B"),
    )
    meta_style = ParagraphStyle(
        "ReportMeta",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9,
        leading=13,
        alignment=2,  # Direita
        textColor=colors.HexColor("#334155"),
    )
    kpi_title_style = ParagraphStyle(
        "KPITitle",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=8,
        leading=10,
        alignment=1,
        textColor=colors.HexColor("#64748B"),
    )
    kpi_val_style = ParagraphStyle(
        "KPIVal",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=13,
        leading=15,
        alignment=1,
        textColor=colors.HexColor("#0F172A"),
    )
    cell_style = ParagraphStyle(
        "TableCell",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=8,
        leading=10,
        textColor=colors.HexColor("#1E293B"),
    )
    cell_bold_style = ParagraphStyle(
        "TableCellBold",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=8,
        leading=10,
        textColor=colors.HexColor("#1E293B"),
    )
    th_style = ParagraphStyle(
        "TableHeader",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=8,
        leading=10,
        textColor=colors.HexColor("#FFFFFF"),
    )
    th_center_style = ParagraphStyle(
        "TableHeaderCenter",
        parent=th_style,
        alignment=1,
    )
    th_right_style = ParagraphStyle(
        "TableHeaderRight",
        parent=th_style,
        alignment=2,
    )
    cell_center_style = ParagraphStyle(
        "TableCellCenter",
        parent=cell_style,
        alignment=1,
    )
    cell_right_style = ParagraphStyle(
        "TableCellRight",
        parent=cell_style,
        alignment=2,
    )

    story = []

    # 1. Cabeçalho Institucional
    now_str = datetime.datetime.now().strftime("%d/%m/%Y às %H:%M")
    header_data = [
        [
            Paragraph("<b>ZAERON CONTÁBIL</b><br/><font color='#C8753D'>Relatório de Conciliação e Classificação</font>", title_style),
            Paragraph(f"<b>{html.escape(office_name)}</b><br/>Responsável: {html.escape(accountant_name)}<br/>Emissão: {now_str}", meta_style),
        ]
    ]
    header_table = Table(header_data, colWidths=[140 * mm, 117 * mm])
    header_table.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 0),
        ("RIGHTPADDING", (0, 0), (-1, -1), 0),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
    ]))
    story.append(header_table)
    story.append(Spacer(1, 8 * mm))

    # 2. Estatísticas do Lote (KPIs)
    total_tx = len(transactions)
    total_saidas = sum(abs(float(t.amount)) for t in transactions if float(t.amount) < 0)
    total_entradas = sum(float(t.amount) for t in transactions if float(t.amount) > 0)
    confident_tx = sum(1 for t in transactions if float(getattr(t, "confidence", 0) or 0) >= 0.8)
    taxa_conciliacao = (confident_tx / total_tx * 100) if total_tx > 0 else 100.0
    confianca_media = (sum(float(getattr(t, "confidence", 1.0) or 1.0) for t in transactions) / total_tx * 100) if total_tx > 0 else 100.0

    kpi_data = [
        [
            Paragraph("TOTAL LANÇAMENTOS", kpi_title_style),
            Paragraph("TOTAL SAÍDAS", kpi_title_style),
            Paragraph("TOTAL ENTRADAS", kpi_title_style),
            Paragraph("TAXA CONCILIAÇÃO", kpi_title_style),
            Paragraph("CONFIANÇA MÉDIA", kpi_title_style),
        ],
        [
            Paragraph(f"{total_tx}", kpi_val_style),
            Paragraph(f"<font color='#DC2626'>R$ {total_saidas:,.2f}</font>", kpi_val_style),
            Paragraph(f"<font color='#059669'>R$ {total_entradas:,.2f}</font>", kpi_val_style),
            Paragraph(f"<font color='#2563EB'>{taxa_conciliacao:.1f}%</font>", kpi_val_style),
            Paragraph(f"<font color='#C8753D'>{confianca_media:.1f}%</font>", kpi_val_style),
        ],
    ]
    kpi_table = Table(kpi_data, colWidths=[51.4 * mm] * 5)
    kpi_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F8FAFC")),
        ("BOX", (0, 0), (-1, -1), 1, colors.HexColor("#E2E8F0")),
        ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
        ("TOPPADDING", (0, 0), (-1, 0), 6),
        ("BOTTOMPADDING", (0, 0), (-1, 0), 2),
        ("TOPPADDING", (0, 1), (-1, 1), 2),
        ("BOTTOMPADDING", (0, 1), (-1, 1), 6),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))
    story.append(kpi_table)
    story.append(Spacer(1, 8 * mm))

    # 3. Tabela de Lançamentos
    table_headers = [
        Paragraph("Data", th_center_style),
        Paragraph("Histórico Contábil", th_style),
        Paragraph("Categoria", th_style),
        Paragraph("Valor (R$)", th_right_style),
        Paragraph("Status", th_center_style),
    ]

    col_widths = [24 * mm, 125 * mm, 46 * mm, 32 * mm, 30 * mm]
    tx_table_data = [table_headers]

    for idx, tx in enumerate(transactions):
        dt_val = tx.date.strftime("%d/%m/%Y") if hasattr(tx.date, "strftime") else str(tx.date)
        if "-" in dt_val and len(dt_val.split("-")) == 3:
            p = dt_val.split("-")
            dt_val = f"{p[2]}/{p[1]}/{p[0]}"

        amt = float(tx.amount)
        amt_str = f"R$ {amt:,.2f}"
        color_tag = "color='#DC2626'" if amt < 0 else "color='#059669'"
        amt_p = Paragraph(f"<font {color_tag}><b>{amt_str}</b></font>", cell_right_style)

        hist_p = Paragraph(html.escape(tx.accounting_history or tx.clean_description or tx.original_description), cell_style)
        cat_p = Paragraph(html.escape(tx.category or "Outras Despesas"), cell_style)
        tx_status = "Conciliado" if float(getattr(tx, "confidence", 0) or 0) >= 0.8 else "Pendente"
        status_p = Paragraph(html.escape(tx_status), cell_center_style)

        tx_table_data.append([
            Paragraph(dt_val, cell_center_style),
            hist_p,
            cat_p,
            amt_p,
            status_p,
        ])

    data_table = Table(tx_table_data, colWidths=col_widths, repeatRows=1)
    data_table_style = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#171B20")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
        ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
    ]

    for row_idx in range(1, len(tx_table_data)):
        if row_idx % 2 == 0:
            data_table_style.append(("BACKGROUND", (0, row_idx), (-1, row_idx), colors.HexColor("#F8FAFC")))

    data_table.setStyle(TableStyle(data_table_style))
    story.append(data_table)

    doc.build(story, canvasmaker=NumberedCanvas)
    return buffer.getvalue()


# ==============================================================================
# FACTORY PRINCIPAL DE EXPORTAÇÃO
# ==============================================================================

def build_export(
    export_format: str,
    transactions: list[Any],
    *,
    office_name: str,
    accountant_name: str,
    source_filename: str = "",
    batch_id: int | str = "",
) -> ExportedFile:
    dataframe = transactions_to_legacy_dataframe(transactions)
    normalized = export_format.lower()
    spreadsheet_safe = dataframe.map(_safe_spreadsheet_text)

    if normalized == "dominio":
        content = gerar_layout_dominio(spreadsheet_safe).encode("utf-8-sig")
        return ExportedFile(content, "text/plain; charset=utf-8", "txt")
    if normalized == "questor":
        content = gerar_layout_questor(spreadsheet_safe).encode("utf-8-sig")
        return ExportedFile(content, "text/plain; charset=utf-8", "txt")
    if normalized == "omie":
        content = gerar_layout_omie(spreadsheet_safe).encode("utf-8-sig")
        return ExportedFile(content, "text/csv; charset=utf-8", "csv")
    if normalized == "fortes":
        content = gerar_layout_fortes_contmatic(spreadsheet_safe).encode("utf-8-sig")
        return ExportedFile(content, "text/csv; charset=utf-8", "csv")
    if normalized in ("xlsx", "excel"):
        content = gerar_excel_zaeron(transactions, office_name, accountant_name)
        return ExportedFile(
            content,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "xlsx",
        )
    if normalized == "csv":
        df_csv = pd.DataFrame(
            [
                {
                    "Data": tx.date.strftime("%d/%m/%Y") if hasattr(tx.date, "strftime") else str(tx.date),
                    "Histórico Contábil": tx.accounting_history or tx.clean_description or tx.original_description,
                    "Categoria": tx.category or "Outras Despesas",
                    "Valor": float(tx.amount),
                    "Status": "Conciliado" if float(getattr(tx, "confidence", 0) or 0) >= 0.8 else "Pendente",
                }
                for tx in transactions
            ]
        )
        safe = df_csv.map(_safe_spreadsheet_text)
        content = safe.to_csv(index=False, sep=";", lineterminator="\r\n").encode(
            "utf-8-sig"
        )
        return ExportedFile(content, "text/csv; charset=utf-8", "csv")
    if normalized == "html":
        safe = _sanitized_html_dataframe(dataframe)
        hours_saved = max(0.2, round(len(dataframe) * 0.025, 1)) if len(dataframe) > 0 else 0.0
        financial_savings = round(hours_saved * 11.0, 2)
        report = gerar_relatorio_gerencial_html(
            safe,
            {
                "escritorio": html.escape(office_name, quote=True),
                "nome": html.escape(accountant_name, quote=True),
                "arquivo_origem": html.escape(source_filename or "Extrato Bancário", quote=True),
                "lote_id": str(batch_id or ""),
                "email": "",
            },
            {
                "horas_salvas": hours_saved,
                "economia_financeira": financial_savings,
            },
        )
        return ExportedFile(report.encode("utf-8"), "text/html; charset=utf-8", "html")
    if normalized == "pdf":
        content = gerar_pdf_zaeron(transactions, office_name, accountant_name)
        return ExportedFile(content, "application/pdf", "pdf")

    raise ValueError(f"Formato de exportação não suportado: {export_format}")
